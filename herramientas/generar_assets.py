#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Genera los recursos de la app «Disponibilidad Casa de Mar» a partir de los archivos de diseño.

Solo hace falta correrlo si cambian las fotos, las siluetas de los .ai o el Excel base.
El uso diario de la app no lo necesita.

Entradas (carpeta --fuentes, con estos nombres):
  Torre_1_Disponibilidad_todos_los_aptos_disponibles.ai   siluetas, etiquetas y niveles de Torre 1
  Torre_2_Disponibilidad_todos_los_aptos_disponibles.ai   lo mismo para Torre 2
  Torre_1_foto_HD.pdf, Torre_2_foto_HD.pdf                fotografías en alta resolución
  CASA_DE_MAR_updated_-_ALAN.xlsx                         inventario (hojas «Torre N Precio Diagrama»)
  DISPONIBILIDAD_URGENCIA_TORRE_2__13_de_oct.pdf          de aquí salen el logo vectorial y la estrella

Salidas (carpeta --salida = raíz del repositorio):
  assets/img/torre-1.jpg, torre-1-preview.jpg, torre-2.jpg, torre-2-preview.jpg
  assets/favicon.svg
  data/geometria.js
  data/inventario.js

Las siluetas se dibujaron sobre una versión de baja resolución de cada foto. La foto HD de
Torre 2 no coincide con esa versión (está escalada ~4 % y desplazada), así que el script registra
la HD contra la original con OpenCV (ECC afín) y la remuestrea para que las siluetas calcen.

Requisitos: pip install pymupdf opencv-python-headless numpy pillow openpyxl
Uso:        python3 herramientas/generar_assets.py --fuentes RUTA --salida .
"""
import argparse
import colorsys
import datetime
import json
import math
import os
import re
import sys

import cv2
import numpy as np
import pymupdf
from openpyxl import load_workbook
from PIL import Image

ARCHIVOS = {
    "T1": {"ai": "Torre_1_Disponibilidad_todos_los_aptos_disponibles.ai", "foto": "Torre_1_foto_HD.pdf"},
    "T2": {"ai": "Torre_2_Disponibilidad_todos_los_aptos_disponibles.ai", "foto": "Torre_2_foto_HD.pdf"},
    "excel": "CASA_DE_MAR_updated_-_ALAN.xlsx",
    "ejemplo": "DISPONIBILIDAD_URGENCIA_TORRE_2__13_de_oct.pdf",
}
HOJAS_EXCEL = {"T1": "Torre 1 Precio Diagrama", "T2": "Torre 2 Precio Diagrama"}
TITULOS = {"T1": "Disponibilidad Torre 01", "T2": "Disponibilidad Torre 02"}
NOMBRES = {"T1": "Torre 1", "T2": "Torre 2"}

# Estilo de las siluetas tomado de los .ai originales (relleno, opacidad del grupo y borde blanco).
ESTILO = {
    "T1": {"verde": "#43A020", "oscuro": "#2D2D2D", "opacidad": 0.75, "borde": 0.5},
    "T2": {"verde": "#43A020", "oscuro": "#1D1D1B", "opacidad": 0.8, "borde": 0.25},
}
VERDE_AI = (0.263, 0.627, 0.125)
ALTO_HD, ALTO_PREVIEW = 2400, 1200
CALIDAD_HD, CALIDAD_PREVIEW = 86, 80


# --------------------------------------------------------------------------- trazados
def fmt(v):
    s = f"{v:.3f}".rstrip("0").rstrip(".")
    return "0" if s in ("-0", "", "-") else s


def subtrazos(items, ox=0.0, oy=0.0):
    """Ítems de PyMuPDF -> lista de subtrazos [('M',p),('L',p),('C',c1,c2,p),('Z',)]."""
    subs, cur = [], None

    def P(p):
        return (p.x - ox, p.y - oy)

    def igual(a, b):
        return abs(a[0] - b[0]) < 1e-3 and abs(a[1] - b[1]) < 1e-3

    for it in items:
        op = it[0]
        if op == "re":
            r = it[1]
            subs.append([("M", P(r.tl)), ("L", P(r.tr)), ("L", P(r.br)), ("L", P(r.bl)), ("Z",)])
            cur = None
        elif op == "qu":
            q = it[1]
            subs.append([("M", P(q.ul)), ("L", P(q.ur)), ("L", P(q.lr)), ("L", P(q.ll)), ("Z",)])
            cur = None
        elif op in ("l", "c"):
            p1 = P(it[1])
            if cur is None or not igual(cur, p1):
                subs.append([("M", p1)])
            if op == "l":
                p2 = P(it[2])
                subs[-1].append(("L", p2))
            else:
                p2 = P(it[4])
                subs[-1].append(("C", P(it[2]), P(it[3]), p2))
            cur = p2
    return subs


def area_firmada(sub):
    pts = [s[-1] for s in sub if s[0] != "Z"]
    return sum(pts[i][0] * pts[(i + 1) % len(pts)][1] - pts[(i + 1) % len(pts)][0] * pts[i][1] for i in range(len(pts))) / 2


def invertir(sub):
    segs = [s for s in sub if s[0] != "Z"]
    cerrado = any(s[0] == "Z" for s in sub)
    prev, tramos = segs[0][1], []
    for s in segs[1:]:
        if s[0] == "L":
            tramos.append(("L", prev, s[1]))
            prev = s[1]
        else:
            tramos.append(("C", prev, s[1], s[2], s[3]))
            prev = s[3]
    out = [("M", prev)]
    for t in reversed(tramos):
        out.append(("L", t[1]) if t[0] == "L" else ("C", t[3], t[2], t[1]))
    if cerrado:
        out.append(("Z",))
    return out


def corregir_huecos(subs):
    """Con relleno 'nonzero' los huecos deben girar al revés que el contorno exterior."""
    if len(subs) < 2:
        return subs
    areas = [area_firmada(s) for s in subs]
    mayor = max(range(len(subs)), key=lambda i: abs(areas[i]))
    if all((a > 0) == (areas[mayor] > 0) for i, a in enumerate(areas) if i != mayor):
        subs = list(subs)
        subs[mayor] = invertir(subs[mayor])
    return subs


def a_svg(subs):
    out = []
    for sub in subs:
        for s in sub:
            if s[0] == "M":
                out.append(f"M{fmt(s[1][0])} {fmt(s[1][1])}")
            elif s[0] == "L":
                out.append(f"L{fmt(s[1][0])} {fmt(s[1][1])}")
            elif s[0] == "C":
                out.append("C" + " ".join(f"{fmt(p[0])} {fmt(p[1])}" for p in s[1:]))
        out.append("Z")
    return "".join(out)


def caja(subs):
    xs = [p[0] for sub in subs for s in sub for p in s[1:]]
    ys = [p[1] for sub in subs for s in sub for p in s[1:]]
    return min(xs), min(ys), max(xs), max(ys)


def aplanar(subs, pasos=8):
    polys = []
    for sub in subs:
        pts, cur = [], None
        for s in sub:
            if s[0] in ("M", "L"):
                cur = s[1]
                pts.append(cur)
            elif s[0] == "C":
                p0, c1, c2, p3 = cur, s[1], s[2], s[3]
                for i in range(1, pasos + 1):
                    t = i / pasos
                    m = 1 - t
                    pts.append((m ** 3 * p0[0] + 3 * m * m * t * c1[0] + 3 * m * t * t * c2[0] + t ** 3 * p3[0],
                                m ** 3 * p0[1] + 3 * m * m * t * c1[1] + 3 * m * t * t * c2[1] + t ** 3 * p3[1]))
                cur = p3
        polys.append(pts)
    return polys


def dentro(pt, polys):
    x, y = pt
    adentro = False
    for poly in polys:
        j = len(poly) - 1
        for i in range(len(poly)):
            xi, yi = poly[i]
            xj, yj = poly[j]
            if (yi > y) != (yj > y) and x < (xj - xi) * (y - yi) / (yj - yi + 1e-12) + xi:
                adentro = not adentro
            j = i
    return adentro


def centroide(polys):
    A = cx = cy = 0.0
    for poly in polys:
        for i in range(len(poly)):
            x0, y0 = poly[i]
            x1, y1 = poly[(i + 1) % len(poly)]
            c = x0 * y1 - x1 * y0
            A += c
            cx += (x0 + x1) * c
            cy += (y0 + y1) * c
    A /= 2
    return (cx / (6 * A), cy / (6 * A)) if abs(A) > 1e-9 else (0, 0)


# --------------------------------------------------------------------------- geometría
def textos(page):
    out = []
    for b in page.get_text("rawdict")["blocks"]:
        for l in b.get("lines", []):
            for s in l["spans"]:
                chars = [c for c in s["chars"] if c["c"].strip()]
                if not chars:
                    continue
                out.append({
                    "texto": re.sub(r"\s+", " ", "".join(c["c"] for c in s["chars"])).strip(),
                    "fuente": s["font"], "tam": s["size"],
                    "x0": min(c["bbox"][0] for c in chars), "x1": max(c["bbox"][2] for c in chars),
                    "y0": min(c["bbox"][1] for c in chars), "y1": max(c["bbox"][3] for c in chars),
                    "base": s["origin"][1],
                })
    return out


def es_verde(c):
    return c is not None and all(abs(a - b) < 0.01 for a, b in zip(c, VERDE_AI))


def es_blanco(c):
    return c is not None and all(abs(a - 1) < 0.01 for a in c)


def clave_nivel(texto):
    t = texto.strip().upper()
    if t == "PH":
        return "PH"
    m = re.search(r"(\d+)", t)
    return str(int(m.group(1))) if m else None


def extraer_geometria(torre, ruta_ai):
    doc = pymupdf.open(ruta_ai)
    page = doc[0]
    dibujos = page.get_drawings()
    formas = []
    for d in dibujos:
        if es_verde(d.get("fill")):
            subs = subtrazos(d["items"])
            polys = aplanar(subs)
            formas.append({"subs": subs, "polys": polys, "c": centroide(polys), "caja": caja(subs)})
    spans = textos(page)
    niveles_txt = [s for s in spans if s["fuente"].startswith("PPWriter")]
    etiquetas = [s for s in spans if not s["fuente"].startswith("PPWriter")]
    for e in etiquetas:
        pt = ((e["x0"] + e["x1"]) / 2, (e["y0"] + e["y1"]) / 2)
        cands = [i for i, f in enumerate(formas) if dentro(pt, f["polys"])] or list(range(len(formas)))
        cands.sort(key=lambda i: math.dist(pt, formas[i]["c"]))
        e["forma"] = cands[0]
    if len({e["forma"] for e in etiquetas}) != len(formas) or len(etiquetas) != len(formas):
        sys.exit(f"[{torre}] No se pudo emparejar cada silueta con su etiqueta ({len(formas)} siluetas, {len(etiquetas)} etiquetas).")
    por_forma = {e["forma"]: e for e in etiquetas}

    asignacion = asignar_t1(formas) if torre == "T1" else asignar_t2(formas, por_forma)

    unidades = {}
    for uid, i in asignacion.items():
        e = por_forma[i]
        unidades[uid] = {
            "d": a_svg(formas[i]["subs"]),
            "etiqueta": {"texto": e["texto"], "cx": round((e["x0"] + e["x1"]) / 2, 2), "base": round(e["base"], 2), "tam": round(e["tam"], 2)},
        }

    # Rótulos «Nv. 0X» con su flecha
    flechas, puntas = [], []
    for d in dibujos:
        r = d["rect"]
        if d["type"] == "s" and es_blanco(d.get("color")) and r.height < 0.5 and r.x1 < 170 and r.width > 15:
            flechas.append({"x1": round(r.x0, 2), "x2": round(r.x1, 2), "y": round((r.y0 + r.y1) / 2, 2), "grosor": round(d.get("width") or 0.25, 3)})
        if d["type"] == "f" and es_blanco(d.get("fill")) and r.x1 < 170 and r.width < 8:
            subs = subtrazos(d["items"])
            puntas.append({"d": a_svg(subs), "y": (caja(subs)[1] + caja(subs)[3]) / 2})
    niveles = []
    for n in sorted(niveles_txt, key=lambda s: s["base"]):
        yc = n["base"] - n["tam"] * 0.3
        fl = min(flechas, key=lambda f: abs(f["y"] - yc)) if flechas else None
        pu = min(puntas, key=lambda p: abs(p["y"] - yc)) if puntas else None
        niveles.append({
            "nivel": clave_nivel(n["texto"]), "texto": n["texto"], "x": round(n["x0"], 2), "base": round(n["base"], 2), "tam": round(n["tam"], 3),
            "flecha": fl if fl and abs(fl["y"] - yc) < 15 else None,
            "punta": pu["d"] if pu and abs(pu["y"] - yc) < 15 else None,
        })
    return {"ancho": page.rect.width, "alto": page.rect.height, "unidades": unidades, "niveles": niveles}


def uid(torre, nivel, pos):
    return f"{torre}-PH{pos}" if nivel == "PH" else f"{torre}-{nivel}{pos:02d}"


def asignar_t1(formas):
    """Torre 1: filas alineadas. De arriba hacia abajo: PH, 9 … 2, 1. En cada fila, izquierda → derecha."""
    orden = sorted(range(len(formas)), key=lambda i: formas[i]["c"][1])
    filas = []
    for i in orden:
        if filas and abs(formas[i]["c"][1] - formas[filas[-1][0]]["c"][1]) < 8:
            filas[-1].append(i)
        else:
            filas.append([i])
    tam = [len(f) for f in filas]
    if tam != [2] + [4] * 8 + [2]:
        sys.exit(f"[T1] Distribución de filas inesperada: {tam}")
    niveles = ["PH"] + [str(n) for n in range(9, 1, -1)] + ["1"]
    res = {}
    for nivel, fila in zip(niveles, filas):
        for pos, i in enumerate(sorted(fila, key=lambda i: formas[i]["c"][0]), 1):
            res[uid("T1", nivel, pos)] = i
    return res


def asignar_t2(formas, por_forma):
    """Torre 2: penthouses en el nivel 9, jardines en el 1 y cuatro columnas escalonadas (niveles 8 → 2)."""
    ph = [i for i in range(len(formas)) if por_forma[i]["texto"].lower().startswith("penthouse")]
    jardin = [i for i in range(len(formas)) if por_forma[i]["texto"].lower().startswith("garden")]
    resto = sorted((i for i in range(len(formas)) if i not in ph and i not in jardin), key=lambda i: formas[i]["c"][0])
    if len(ph) != 4 or len(jardin) != 2 or len(resto) != 28:
        sys.exit(f"[T2] Conteo inesperado: {len(ph)} PH, {len(jardin)} jardines, {len(resto)} intermedias")
    columnas = [resto[k * 7:(k + 1) * 7] for k in range(4)]
    for a, b in zip(columnas, columnas[1:]):
        if max(formas[i]["c"][0] for i in a) >= min(formas[i]["c"][0] for i in b):
            sys.exit("[T2] Las columnas de siluetas se solapan; revisar el .ai")
    res = {}
    for pos, i in enumerate(sorted(ph, key=lambda i: formas[i]["c"][0]), 1):
        res[uid("T2", "9", pos)] = i
    for pos, i in enumerate(sorted(jardin, key=lambda i: formas[i]["c"][0]), 1):
        res[uid("T2", "1", pos)] = i
    for pos, col in enumerate(columnas, 1):
        for nivel, i in zip(range(8, 1, -1), sorted(col, key=lambda i: formas[i]["c"][1])):
            res[uid("T2", str(nivel), pos)] = i
    return res


# --------------------------------------------------------------------------- fotos
def imagen_mayor(page):
    infos = page.get_image_info(xrefs=True)
    return max(infos, key=lambda i: i["width"] * i["height"])


def pixeles(doc, xref):
    """RGB uint8; si la imagen trae máscara suave se compone sobre blanco (como en el PDF)."""
    pix = pymupdf.Pixmap(doc, xref)
    if pix.alpha:
        pix = pymupdf.Pixmap(pix, 0)
    if pix.n != 3:
        pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
    rgb = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3).astype(np.float32)
    smask = doc.extract_image(xref).get("smask", 0)
    if smask:
        m = pymupdf.Pixmap(doc, smask)
        a = np.frombuffer(m.samples, dtype=np.uint8).reshape(m.height, m.width, m.n)[:, :, 0]
        if a.shape != rgb.shape[:2]:
            a = cv2.resize(a, (rgb.shape[1], rgb.shape[0]), interpolation=cv2.INTER_LINEAR)
        a = a.astype(np.float32)[:, :, None] / 255.0
        rgb = rgb * a + 255.0 * (1 - a)
    return np.clip(rgb, 0, 255).astype(np.uint8)


def en_pagina(arr, tr, W, H, S):
    a, _, _, d, e, f = tr
    h, w = arr.shape[:2]
    M = np.array([[a / w * S, 0, e * S], [0, d / h * S, f * S]], dtype=np.float64)
    return cv2.warpAffine(arr, M, (int(round(W * S)), int(round(H * S))), flags=cv2.INTER_AREA, borderMode=cv2.BORDER_REPLICATE)


def generar_foto(torre, ruta_ai, ruta_hd, salida):
    ai = pymupdf.open(ruta_ai)
    pa = ai[0]
    hd = pymupdf.open(ruta_hd)
    ph = hd[0]
    W, H = pa.rect.width, pa.rect.height
    if abs(ph.rect.width - W) > 0.05 or abs(ph.rect.height - H) > 0.05:
        sys.exit(f"[{torre}] La foto HD y el .ai no tienen el mismo tamaño de mesa de trabajo")
    ref_info, hd_info = imagen_mayor(pa), imagen_mayor(ph)
    ref_px, hd_px = pixeles(ai, ref_info["xref"]), pixeles(hd, hd_info["xref"])

    # 1) Registro: ¿dónde cae cada punto de la mesa de trabajo dentro de la foto HD?
    S = 2.0
    g_ref = cv2.cvtColor(en_pagina(ref_px, ref_info["transform"], W, H, S), cv2.COLOR_RGB2GRAY).astype(np.float32)
    g_hd = cv2.cvtColor(en_pagina(hd_px, hd_info["transform"], W, H, S), cv2.COLOR_RGB2GRAY).astype(np.float32)
    warp = np.eye(2, 3, dtype=np.float32)
    cc, warp = cv2.findTransformECC(g_ref, g_hd, warp, cv2.MOTION_AFFINE,
                                    (cv2.TERM_CRITERIA_EPS | cv2.TERM_CRITERIA_COUNT, 300, 1e-7), None, 5)
    A = warp[:, :2].astype(np.float64)
    T = warp[:, 2].astype(np.float64) / S        # traslación en puntos
    print(f"  [{torre}] registro HD↔original: correlación {cc:.4f}, escala {A[0, 0]:.4f}/{A[1, 1]:.4f}, desplazamiento {T[0]:.2f}, {T[1]:.2f} pt")

    # 2) Zona de la mesa de trabajo que la foto HD cubre de verdad
    a, _, _, d, e, f = hd_info["transform"]
    ancho_vis = min(W, (e + a - T[0] - max(0.0, A[0, 1] * H)) / A[0, 0] - 0.25)
    alto_vis = min(H, (f + d - T[1] - max(0.0, A[1, 0] * W)) / A[1, 1] - 0.25)
    if (e - T[0]) / A[0, 0] > 0.01 or (f - T[1]) / A[1, 1] > 0.01:
        sys.exit(f"[{torre}] La foto HD no cubre la esquina superior izquierda de la mesa de trabajo")
    ancho_vis, alto_vis = round(ancho_vis, 3), round(alto_vis, 3)

    # 3) Remuestreo a densidad nativa y reducción final con filtro de área
    h_src, w_src = hd_px.shape[:2]
    R = (w_src / a) * A[0, 0]                     # píxeles HD por punto de la mesa de trabajo
    out_w, out_h = int(round(ancho_vis * R)), int(round(alto_vis * R))
    sx, sy = w_src / a, h_src / d
    M = np.array([
        [A[0, 0] * sx / R, A[0, 1] * sx / R, (T[0] - e) * sx],
        [A[1, 0] * sy / R, A[1, 1] * sy / R, (T[1] - f) * sy],
    ], dtype=np.float64)
    nativo = cv2.warpAffine(hd_px, M, (out_w, out_h), flags=cv2.INTER_LANCZOS4 | cv2.WARP_INVERSE_MAP, borderMode=cv2.BORDER_REPLICATE)
    os.makedirs(os.path.join(salida, "assets", "img"), exist_ok=True)
    n = torre[1]
    for alto, calidad, sufijo in ((ALTO_HD, CALIDAD_HD, ""), (ALTO_PREVIEW, CALIDAD_PREVIEW, "-preview")):
        ancho = int(round(alto * ancho_vis / alto_vis))
        img = cv2.resize(nativo, (ancho, alto), interpolation=cv2.INTER_AREA)
        ruta = os.path.join(salida, "assets", "img", f"torre-{n}{sufijo}.jpg")
        Image.fromarray(img).save(ruta, "JPEG", quality=calidad, optimize=True, progressive=True)
        print(f"  [{torre}] {ruta}: {ancho}×{alto}px, {os.path.getsize(ruta) / 1024:.0f} KB")
    return ancho_vis, alto_vis


# --------------------------------------------------------------------------- logo y estrella
def extraer_logo_y_estrella(ruta_pdf):
    page = pymupdf.open(ruta_pdf)[0]
    dibujos = page.get_drawings()
    region = pymupdf.Rect(850, 28, 1090, 80)
    logo = [d for d in dibujos if region.contains(d["rect"]) and d.get("fill")]
    x0 = min(d["rect"].x0 for d in logo)
    y0 = min(d["rect"].y0 for d in logo)
    x1 = max(d["rect"].x1 for d in logo)
    y1 = max(d["rect"].y1 for d in logo)
    trazos = [a_svg(corregir_huecos(subtrazos(d["items"], x0, y0))) for d in logo]
    icono = next(a_svg(corregir_huecos(subtrazos(d["items"], d["rect"].x0, d["rect"].y0))) for d in logo if abs(d["rect"].x0 - x0) < 0.01)
    lado_icono = next(d["rect"].width for d in logo if abs(d["rect"].x0 - x0) < 0.01)

    estrellas = [d for d in dibujos if d.get("fill") and abs(d["fill"][0] - 0.867) < 0.02 and abs(d["fill"][1] - 0.773) < 0.02]
    est = min(estrellas, key=lambda d: d["rect"].width)
    r = est["rect"]
    subs = subtrazos(est["items"], r.x0, r.y0)
    escala = 1 / r.width
    subs = [[(s[0],) + tuple((p[0] * escala, p[1] * escala) for p in s[1:]) for s in sub] for sub in subs]
    return (
        {"ancho": round(x1 - x0, 3), "alto": round(y1 - y0, 3), "d": "".join(trazos)},
        {"icono": icono, "lado": round(lado_icono, 3)},
        {"ancho": 1, "alto": round(r.height / r.width, 4), "d": a_svg(subs), "color": "#DDC517"},
    )


# --------------------------------------------------------------------------- inventario (Excel)
TEMA_OFFICE = ["FFFFFF", "000000", "E7E6E6", "44546A", "4472C4", "ED7D31", "A5A5A5", "FFC000", "5B9BD5", "70AD47"]


def con_tinte(hexa, tinte):
    r, g, b = (int(hexa[i:i + 2], 16) / 255 for i in (0, 2, 4))
    h, l, s = colorsys.rgb_to_hls(r, g, b)
    l = l * (1 + tinte) if tinte < 0 else l * (1 - tinte) + tinte
    r, g, b = colorsys.hls_to_rgb(h, l, s)
    return "".join(f"{round(v * 255):02X}" for v in (r, g, b))


def color_relleno(celda):
    f = celda.fill
    if f is None or f.fill_type in (None, "none"):
        return None
    c = f.fgColor
    if c.type == "rgb" and c.rgb and isinstance(c.rgb, str):
        return c.rgb[-6:].upper()
    if c.type == "theme" and c.theme is not None:
        return con_tinte(TEMA_OFFICE[c.theme] if c.theme < len(TEMA_OFFICE) else "FFFFFF", c.tint or 0)
    return "808080"


def es_disponible(hexa):
    """Sin color, blanco o azul claro (DISPONIBLE en la leyenda) = disponible. Gris, amarillo, etc. = reservado."""
    if hexa is None:
        return True
    r, g, b = (int(hexa[i:i + 2], 16) for i in (0, 2, 4))
    if min(r, g, b) >= 245:
        return True
    h, l, s = colorsys.rgb_to_hls(r / 255, g / 255, b / 255)
    return 190 <= h * 360 <= 230 and l >= 0.75 and s >= 0.2


def decimal(txt):
    txt = txt.strip()
    if "," in txt and "." in txt:
        txt = txt.replace(".", "").replace(",", ".") if txt.rfind(",") > txt.rfind(".") else txt.replace(",", "")
    elif "," in txt:
        txt = txt.replace(",", ".") if len(txt.split(",")[-1]) != 3 else txt.replace(",", "")
    return float(txt)


def leer_celda(texto):
    lineas = [l.strip() for l in str(texto).replace("\r", "").split("\n") if l.strip()]
    primera = lineas[0]
    modelo = re.sub(r"\s*-\s*c\s*/\s*jacuzzi.*$", "", primera, flags=re.I).strip()
    area = re.search(r"([\d.,]+)\s*m²", texto)
    hab = re.search(r"(\d+)\s*Hab", texto, re.I)
    precio = re.search(r"\$\s*([\d.,]+)", texto)
    return {
        "modelo": re.sub(r"\s+", " ", modelo),
        "jacuzzi": bool(re.search(r"jacuzzi", texto, re.I)),
        "area": round(decimal(area.group(1)), 2) if area else None,
        "hab": int(hab.group(1)) if hab else None,
        "precio": int(re.sub(r"\D", "", precio.group(1))) if precio and re.sub(r"\D", "", precio.group(1)) else None,
    }


def nombre_lista(modelo):
    m = re.match(r"Penthouse\s*-\s*(Zenith\s*\+?)\s*$", modelo, re.I)
    return re.sub(r"\s+", " ", m.group(1)).strip() if m else modelo


def leer_inventario(ruta_xlsx):
    wb = load_workbook(ruta_xlsx)
    inventario = {}
    for torre, hoja in HOJAS_EXCEL.items():
        ws = wb[hoja]
        unidades, niveles = [], []
        for fila in ws.iter_rows(min_row=1, max_row=ws.max_row):
            etiqueta = str(fila[0].value or "").strip()
            if etiqueta.upper() == "PH":
                nivel = "PH"
            elif re.match(r"^piso\s*\d+$", etiqueta, re.I):
                nivel = str(int(re.search(r"\d+", etiqueta).group()))
            else:
                continue
            celdas = [c for c in fila[1:5] if type(c).__name__ != "MergedCell"]
            pos = 0
            for c in celdas:
                pos += 1
                if c.value is None or not str(c.value).strip():
                    continue
                datos = leer_celda(str(c.value))
                unidades.append({
                    "id": uid(torre, nivel, pos), "nivel": nivel, "pos": pos,
                    "modelo": datos["modelo"], "nombre": nombre_lista(datos["modelo"]),
                    "area": datos["area"], "hab": datos["hab"], "jacuzzi": datos["jacuzzi"], "precio": datos["precio"],
                    "disponible": es_disponible(color_relleno(c)),
                })
            niveles.append(nivel)
        inventario[torre] = {"niveles": niveles, "unidades": unidades}
    return inventario


# --------------------------------------------------------------------------- principal
def escribir_js(ruta, nombre, valor):
    with open(ruta, "w", encoding="utf-8") as fh:
        fh.write("/* Generado por herramientas/generar_assets.py. No editar a mano: vuelve a correr el script. */\n")
        fh.write("window.CDM = window.CDM || {};\n")
        fh.write(f"CDM.{nombre} = ")
        json.dump(valor, fh, ensure_ascii=False, separators=(",", ":"))
        fh.write(";\n")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--fuentes", required=True, help="carpeta con los .ai, las fotos HD, el Excel y el PDF de ejemplo")
    ap.add_argument("--salida", default=".", help="raíz del repositorio de la app")
    args = ap.parse_args()
    F = lambda nombre: os.path.join(args.fuentes, nombre)

    print("Inventario desde el Excel…")
    inv = leer_inventario(F(ARCHIVOS["excel"]))
    print("Logo y estrella…")
    logo, icono, estrella = extraer_logo_y_estrella(F(ARCHIVOS["ejemplo"]))

    geometria = {"torres": {}, "logo": logo, "estrella": estrella}
    inventario = {"version": datetime.date.today().isoformat(), "fuente": ARCHIVOS["excel"], "torres": {}}
    for torre in ("T1", "T2"):
        print(f"{NOMBRES[torre]}: siluetas y foto…")
        geo = extraer_geometria(torre, F(ARCHIVOS[torre]["ai"]))
        ancho_vis, alto_vis = generar_foto(torre, F(ARCHIVOS[torre]["ai"]), F(ARCHIVOS[torre]["foto"]), args.salida)
        datos = inv[torre]
        ids_excel = {u["id"] for u in datos["unidades"]}
        if ids_excel != set(geo["unidades"]):
            sys.exit(f"[{torre}] Las unidades del Excel y las siluetas no coinciden: {sorted(ids_excel ^ set(geo['unidades']))}")
        avisos = 0
        for u in datos["unidades"]:
            etiqueta = geo["unidades"][u["id"]]["etiqueta"]["texto"]
            if etiqueta.lower() != u["modelo"].lower():
                avisos += 1
                print(f"  aviso {u['id']}: la foto dice «{etiqueta}» y el Excel «{u['modelo']}»")
        print(f"  {len(datos['unidades'])} unidades, {sum(u['disponible'] for u in datos['unidades'])} disponibles según el Excel, {avisos} avisos de nombre")
        geometria["torres"][torre] = {
            "ancho": round(ancho_vis, 3), "alto": round(alto_vis, 3),
            "foto": f"assets/img/torre-{torre[1]}.jpg", "fotoPreview": f"assets/img/torre-{torre[1]}-preview.jpg",
            "estilo": ESTILO[torre], "niveles": geo["niveles"], "unidades": geo["unidades"],
        }
        inventario["torres"][torre] = {
            "id": torre, "nombre": NOMBRES[torre], "titulo": TITULOS[torre],
            "niveles": [{"clave": n, "nombre": f"Nivel {n}"} for n in datos["niveles"]],
            "unidades": datos["unidades"],
        }

    os.makedirs(os.path.join(args.salida, "data"), exist_ok=True)
    escribir_js(os.path.join(args.salida, "data", "geometria.js"), "geometria", geometria)
    escribir_js(os.path.join(args.salida, "data", "inventario.js"), "inventario", inventario)
    lado = icono["lado"]
    with open(os.path.join(args.salida, "assets", "favicon.svg"), "w", encoding="utf-8") as fh:
        fh.write(f'<svg xmlns="http://www.w3.org/2000/svg" viewBox="-2 -2 {fmt(lado + 4)} {fmt(lado + 4)}">'
                 f'<rect x="-2" y="-2" width="{fmt(lado + 4)}" height="{fmt(lado + 4)}" fill="#fff"/><path d="{icono["icono"]}"/></svg>')
    print("Listo.")


if __name__ == "__main__":
    main()
